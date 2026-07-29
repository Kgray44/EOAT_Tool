from __future__ import annotations

import csv
import json
import mimetypes
import re
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any
from uuid import NAMESPACE_URL, uuid4, uuid5

from openpyxl import load_workbook
from sqlalchemy import delete, select, text
from sqlalchemy.orm import Session

from server.eoat_api.database import models as db
from server.eoat_api.database.session import create_session_factory
from server.eoat_api.release_provenance import ensure_application_release
from tools.eoat_location_normalization import load_policy, normalized_source_rows
from tools.migration.excel_to_mysql import MISSING_TOKENS, _checksum, _rows, _text

SCHEMA_REVISION = "20260729_0009"
MACHINE_PATTERN = re.compile(r"^\d+$")
RESOLUTION_STATUSES = {"UNRESOLVED", "DEFERRED", "RESOLVED", "NOT_APPLICABLE", "REJECTED_WITH_REASON"}


@dataclass
class ReviewItem:
    issue_id: str
    severity: str
    issue_code: str
    source_workbook: str
    source_sheet: str
    source_row: int | None
    source_identifier: str
    affected_entity_type: str
    affected_entity_identifier: str
    field_name: str
    source_value: str
    conflicting_values: list[str] = field(default_factory=list)
    related_source_rows: list[int] = field(default_factory=list)
    current_proposed_action: str = ""
    allowed_resolution_actions: list[str] = field(default_factory=list)
    resolution_status: str = "UNRESOLVED"
    resolution_value: str = ""
    resolution_reason: str = ""
    resolved_by: str = ""
    resolved_at: str = ""


@dataclass
class ImportResult:
    source_workbook: str
    source_checksum: str
    import_batch_uuid: str
    schema_revision: str = SCHEMA_REVISION
    started_at: str = ""
    completed_at: str = ""
    rows_discovered_by_sheet: dict[str, int] = field(default_factory=dict)
    records_imported_by_table: dict[str, int] = field(default_factory=dict)
    records_skipped: int = 0
    records_rejected: int = 0
    warnings: int = 0
    unresolved_issues: int = 0
    resolved_issues: int = 0
    duplicate_identifiers: dict[str, int] = field(default_factory=dict)
    missing_relationships: int = 0
    photo_path_validation: dict[str, int] = field(default_factory=dict)
    document_path_validation: dict[str, int] = field(default_factory=dict)
    foreign_key_validation: str = "NOT_RUN"
    transaction_result: str = "NOT_STARTED"
    source_workbook_unchanged: bool = False
    deferred_parts: int = 0
    deferred_installations: int = 0
    already_imported: bool = False
    issues: list[ReviewItem] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _slug(value: Any) -> str:
    text_value = re.sub(r"[^a-z0-9]+", "_", _text(value).casefold()).strip("_")
    return text_value[:64] or "unknown"


def _missing(value: Any) -> bool:
    return _text(value).casefold() in MISSING_TOKENS


def _integer(value: Any) -> int | None:
    match = re.search(r"-?\d+", _text(value))
    return int(match.group()) if match else None


def _boolean(value: Any) -> bool | None:
    normalized = _text(value).casefold()
    if normalized in {"yes", "y", "true", "present", "1"}:
        return True
    if normalized in {"no", "n", "false", "not present", "0"}:
        return False
    return None


def _datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    text_value = _text(value)
    if not text_value:
        return None
    try:
        return datetime.fromisoformat(text_value)
    except ValueError:
        return None


def _json_safe(values: dict[str, Any]) -> dict[str, Any]:
    return {key: value.isoformat() if isinstance(value, date | datetime) else value for key, value in values.items()}


def _issue(
    source: Path,
    code: str,
    sheet: str,
    row: int | None,
    source_identifier: str,
    entity_type: str,
    entity_identifier: str,
    field_name: str,
    value: Any,
    action: str,
    *,
    conflicting: list[str] | None = None,
    related_rows: list[int] | None = None,
    status: str = "UNRESOLVED",
    severity: str = "WARNING",
) -> ReviewItem:
    stable = "|".join([str(source), sheet, str(row or 0), code, source_identifier, field_name, _text(value)])
    return ReviewItem(
        issue_id=str(uuid5(NAMESPACE_URL, stable)),
        severity=severity,
        issue_code=code,
        source_workbook=str(source),
        source_sheet=sheet,
        source_row=row,
        source_identifier=source_identifier,
        affected_entity_type=entity_type,
        affected_entity_identifier=entity_identifier,
        field_name=field_name,
        source_value=_text(value),
        conflicting_values=conflicting or [],
        related_source_rows=related_rows or [],
        current_proposed_action=action,
        allowed_resolution_actions=["DEFER", "SUPPLY_VERIFIED_VALUE", "MARK_NOT_APPLICABLE", "REJECT_WITH_REASON"],
        resolution_status=status,
    )


def build_review_items(
    source_workbook: str | Path,
) -> tuple[list[ReviewItem], dict[str, list[tuple[int, dict[str, Any]]]]]:
    source = Path(source_workbook).resolve()
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows = {name: _rows(workbook[name]) for name in ("EOAT Inventory", "Photo Index")}
    workbook.close()
    items: list[ReviewItem] = []
    eoat_rows: dict[str, list[tuple[int, dict[str, Any]]]] = defaultdict(list)
    audited_locations: dict[str, dict[str, list[int]]] = defaultdict(lambda: defaultdict(list))
    part_candidates: dict[tuple[str, str], list[int]] = defaultdict(list)

    for row_number, row in rows["EOAT Inventory"]:
        eoat_id = _text(row.get("EOAT Assembly ID"))
        audit_id = _text(row.get("Audit ID"))
        machine = _text(row.get("Press/Machine #"))
        tool = _text(row.get("Tool #"))
        eoat_rows[eoat_id].append((row_number, row))
        if _missing(machine):
            items.append(
                _issue(
                    source,
                    "MISSING_MACHINE",
                    "EOAT Inventory",
                    row_number,
                    audit_id,
                    "EOAT",
                    eoat_id,
                    "Press/Machine #",
                    machine,
                    "Import the audit and EOAT; defer the machine relationship.",
                )
            )
        elif not MACHINE_PATTERN.fullmatch(machine):
            items.append(
                _issue(
                    source,
                    "AMBIGUOUS_MACHINE_VALUE",
                    "EOAT Inventory",
                    row_number,
                    audit_id,
                    "Machine",
                    machine,
                    "Press/Machine #",
                    machine,
                    "Preserve the value in traceability; do not create a machine relationship.",
                )
            )
        elif _text(row.get("Entry Type")).casefold() == "audited":
            audited_locations[eoat_id][machine].append(row_number)
        if _missing(tool):
            items.append(
                _issue(
                    source,
                    "MISSING_TOOL",
                    "EOAT Inventory",
                    row_number,
                    audit_id,
                    "EOAT",
                    eoat_id,
                    "Tool #",
                    tool,
                    "Import the audit and EOAT; defer the tool relationship.",
                )
            )
        part_name = _text(row.get("Part Name/Description"))
        if part_name:
            part_candidates[(tool, part_name)].append(row_number)

    stable_fields = ("EOAT Type", "Connection Type", "Cleanroom/Non-Cleanroom")
    for eoat_id, grouped in eoat_rows.items():
        for field_name in stable_fields:
            populated = [(number, _text(row.get(field_name))) for number, row in grouped if _text(row.get(field_name))]
            values = sorted({value for _, value in populated})
            if len(values) > 1:
                baseline = populated[0][1]
                for row_number, value in populated[1:]:
                    if value != baseline:
                        items.append(
                            _issue(
                                source,
                                "CONFLICTING_EOAT_ATTRIBUTE",
                                "EOAT Inventory",
                                row_number,
                                eoat_id,
                                "EOAT",
                                eoat_id,
                                field_name,
                                value,
                                "Leave the normalized field unknown until an administrator resolves it.",
                                conflicting=values,
                                related_rows=[number for number, _ in grouped],
                            )
                        )
    for eoat_id, locations in audited_locations.items():
        if len(locations) > 1:
            related = sorted(number for numbers in locations.values() for number in numbers)
            items.append(
                _issue(
                    source,
                    "CONFLICTING_CURRENT_ASSIGNMENT",
                    "EOAT Inventory",
                    None,
                    eoat_id,
                    "EOAT",
                    eoat_id,
                    "Press/Machine #",
                    " | ".join(sorted(locations)),
                    "Do not create an active installation or storage assignment.",
                    conflicting=sorted(locations),
                    related_rows=related,
                )
            )
    for (tool, part_name), source_rows in sorted(part_candidates.items(), key=lambda value: value[1][0]):
        row_number = source_rows[0]
        items.append(
            _issue(
                source,
                "POSSIBLE_PART_NOT_CONFIRMED",
                "EOAT Inventory",
                row_number,
                tool,
                "Part",
                part_name,
                "Part Name/Description",
                part_name,
                "Preserve the candidate; do not create a part or tool-part relationship.",
                related_rows=source_rows,
                status="DEFERRED",
            )
        )
    for eoat_id, grouped in sorted(eoat_rows.items()):
        if eoat_id:
            related = [number for number, _ in grouped]
            items.append(
                _issue(
                    source,
                    "INSTALLATION_DATE_UNKNOWN",
                    "EOAT Inventory",
                    None,
                    eoat_id,
                    "EOATInstallation",
                    eoat_id,
                    "installed_at",
                    "",
                    "Do not fabricate installation history.",
                    related_rows=related,
                    status="DEFERRED",
                )
            )
            items.append(
                _issue(
                    source,
                    "CURRENT_LOCATION_UNKNOWN",
                    "EOAT Inventory",
                    None,
                    eoat_id,
                    "EOAT",
                    eoat_id,
                    "current_location",
                    "",
                    "Return Unknown / Not Verified until authoritative evidence exists.",
                    related_rows=related,
                    status="DEFERRED",
                )
            )
    for row_number, row in rows["Photo Index"]:
        relative = _text(row.get("Stored Relative Path"))
        folder = _text(row.get("Folder Path"))
        filename = _text(row.get("Stored Filename")) or _text(row.get("Photo Filename"))
        if not relative and not (folder and filename):
            photo_id = _text(row.get("Photo ID"))
            items.append(
                _issue(
                    source,
                    "PLACEHOLDER_PHOTO_ROW",
                    "Photo Index",
                    row_number,
                    photo_id,
                    "Photo",
                    photo_id,
                    "Stored Relative Path",
                    "",
                    "Preserve the import row but do not create document/photo records.",
                    status="DEFERRED",
                )
            )
    return items, rows


def write_review_reports(items: list[ReviewItem], output_directory: str | Path) -> tuple[Path, Path, Path]:
    output = Path(output_directory)
    output.mkdir(parents=True, exist_ok=True)
    json_path = output / "import_review_items.json"
    csv_path = output / "import_review_items.csv"
    md_path = output / "import_review_summary.md"
    records = [asdict(item) for item in items]
    json_path.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]) if records else list(ReviewItem.__annotations__))
        writer.writeheader()
        for record in records:
            writer.writerow(
                {
                    key: json.dumps(value, ensure_ascii=False) if isinstance(value, list) else value
                    for key, value in record.items()
                }
            )
    counts: dict[str, int] = defaultdict(int)
    for item in items:
        counts[item.issue_code] += 1
    table = "\n".join(f"| {code} | {count} |" for code, count in sorted(counts.items()))
    md_path.write_text(
        "# MySQL Import Review Summary\n\n"
        f"- Review items: **{len(items)}**\n"
        f"- Unresolved: **{sum(item.resolution_status == 'UNRESOLVED' for item in items)}**\n"
        f"- Deferred: **{sum(item.resolution_status == 'DEFERRED' for item in items)}**\n\n"
        "| Issue code | Count |\n|---|---:|\n" + table + "\n\n"
        "Items are stable-ID records. Values remain unresolved unless an administrator supplies evidence and records a reason.\n",
        encoding="utf-8",
    )
    return md_path, json_path, csv_path


def _lookup(session: Session, model: type, value: Any, *, display: str | None = None) -> int | None:
    if _missing(value):
        return None
    code = _slug(value)
    record = session.scalar(select(model).where(model.code == code))
    if record is None:
        record = model(code=code, display_name=display or _text(value), description="Imported lookup value")
        session.add(record)
        session.flush()
    return record.id


def _reset_imported_data(session: Session) -> None:
    ordered = [
        db.ChangeFeed,
        db.EntityHistoryEvent,
        db.DocumentLink,
        db.Photo,
        db.Document,
        db.AuditRecord,
        db.EOATMachineCompatibility,
        db.EOATToolCompatibility,
        db.ToolMachineCompatibility,
        db.EOATInstallation,
        db.EOATStorageAssignment,
        db.ToolPart,
        db.Part,
        db.Tool,
        db.MachineRobotAssignment,
        db.Robot,
        db.Machine,
        db.EOAT,
        db.StorageLocation,
        db.Area,
        db.Plant,
        db.ImportIssue,
        db.ImportRow,
        db.ImportBatch,
    ]
    for model in ordered:
        session.execute(delete(model))


def _canonical_profile(grouped: list[tuple[int, dict[str, Any]]], field_name: str) -> str | None:
    values = {_text(row.get(field_name)) for _, row in grouped if _text(row.get(field_name))}
    return next(iter(values)) if len(values) == 1 else None


def execute_import(
    source_workbook: str | Path,
    *,
    batch_name: str,
    reset_imported_data: bool = False,
    report_output: str | Path = "reports/mysql_import",
) -> ImportResult:
    source = Path(source_workbook).resolve()
    before_stat = source.stat()
    before_checksum = _checksum(source)
    started = datetime.now(timezone.utc)
    review_items, rows = build_review_items(source)
    write_review_reports(review_items, report_output)
    result = ImportResult(
        source_workbook=str(source),
        source_checksum=before_checksum,
        import_batch_uuid=str(uuid4()),
        started_at=started.isoformat(),
        rows_discovered_by_sheet={name: len(values) for name, values in rows.items()},
        issues=review_items,
    )
    factory = create_session_factory(migration=True)
    with factory() as session:
        existing = session.scalar(
            select(db.ImportBatch).where(
                db.ImportBatch.source_file_checksum == before_checksum,
                db.ImportBatch.status == "COMPLETED",
                db.ImportBatch.dry_run.is_(False),
            )
        )
        session.rollback()
        if existing is not None and not reset_imported_data:
            result.import_batch_uuid = existing.batch_uuid
            result.already_imported = True
            result.transaction_result = "SAFE_STOP_ALREADY_IMPORTED"
            result.records_skipped = sum(result.rows_discovered_by_sheet.values())
        else:
            with session.begin():
                if reset_imported_data:
                    _reset_imported_data(session)
                batch = db.ImportBatch(
                    batch_uuid=result.import_batch_uuid,
                    batch_name=batch_name,
                    source_type="excel_workbook",
                    source_file_name=str(source),
                    source_file_checksum=before_checksum,
                    started_at=started,
                    status="RUNNING",
                    dry_run=False,
                    application_release_id=ensure_application_release(session).id,
                    records_discovered=sum(result.rows_discovered_by_sheet.values()),
                    notes=f"Controlled supported-record import at schema {SCHEMA_REVISION}",
                )
                session.add(batch)
                session.flush()
                plant = db.Plant(
                    plant_code="P4", plant_name="Plant 4", source_system="legacy_excel", source_import_batch_id=batch.id
                )
                session.add(plant)
                session.flush()
                area_map: dict[str, db.Area] = {}
                for _, row in rows["EOAT Inventory"]:
                    label = _text(row.get("Plant/Area")) or "Unknown"
                    if label not in area_map:
                        classification = _lookup(
                            session, db.CleanroomClassification, row.get("Cleanroom/Non-Cleanroom")
                        )
                        area = db.Area(
                            plant_id=plant.id,
                            area_code=_slug(label),
                            area_name=label,
                            cleanroom_classification_id=classification,
                            source_system="legacy_excel",
                            source_import_batch_id=batch.id,
                        )
                        session.add(area)
                        session.flush()
                        area_map[label] = area
                # Resolve audited evidence to governed physical identities before
                # grouping.  Compatibility-only rows remain provenance and can
                # never cause a physical EOAT to be created.
                raw_inventory_rows = list(rows["EOAT Inventory"])
                normalized_inventory = normalized_source_rows(dict(raw_inventory_rows))
                rows["EOAT Inventory"] = [
                    (number, normalized_inventory[number]) for number, _ in raw_inventory_rows
                ]
                split_source_identifiers = {
                    str(split["source_identifier"])
                    for split in load_policy()["physical_unit_splits"]
                }
                inventory_groups: dict[str, list[tuple[int, dict[str, Any]]]] = defaultdict(list)
                for number, row in rows["EOAT Inventory"]:
                    if _text(row.get("Entry Type")).casefold() == "audited":
                        inventory_groups[_text(row.get("EOAT Assembly ID"))].append((number, row))
                eoat_map: dict[str, db.EOAT] = {}
                for identifier, grouped in sorted(inventory_groups.items()):
                    if not identifier:
                        continue
                    first = grouped[0][1]
                    eoat_type = _canonical_profile(grouped, "EOAT Type")
                    connection = _canonical_profile(grouped, "Connection Type")
                    cleanroom = _canonical_profile(grouped, "Cleanroom/Non-Cleanroom")
                    eoat = db.EOAT(
                        business_identifier=identifier,
                        display_name=identifier,
                        description=_text(first.get("Part Name/Description")) or None,
                        eoat_type_id=_lookup(session, db.EOATType, eoat_type),
                        connection_type_id=_lookup(session, db.ConnectionType, connection),
                        cleanroom_classification_id=_lookup(session, db.CleanroomClassification, cleanroom),
                        status_id=_lookup(session, db.AssetStatus, first.get("Status")),
                        number_of_parts_picked=_integer(first.get("Number of Parts Picked")),
                        number_of_vacuum_cups=_integer(first.get("# of Cups")),
                        number_of_grippers=_integer(first.get("# of Grippers")),
                        vacuum_present=_boolean(first.get("Vacuum Confirmation Present?")),
                        sensors_present=_boolean(first.get("Sensors Present?")),
                        part_present_sensor_present=_boolean(first.get("Part-Present Detection Present?")),
                        vacuum_confirmation_sensor_present=_boolean(first.get("Vacuum Confirmation Present?")),
                        quick_disconnect_present=_boolean(first.get("Quick Disconnects Present?")),
                        cup_material=_text(first.get("Cup Type/Material")) or None,
                        notes=_text(first.get("Notes")) or None,
                        source_system="legacy_excel",
                        source_import_batch_id=batch.id,
                    )
                    session.add(eoat)
                    session.flush()
                    eoat_map[identifier] = eoat
                machine_map: dict[str, db.Machine] = {}
                tool_map: dict[str, db.Tool] = {}
                for _, row in rows["EOAT Inventory"]:
                    machine_number = _text(row.get("Press/Machine #"))
                    area = area_map[_text(row.get("Plant/Area")) or "Unknown"]
                    if MACHINE_PATTERN.fullmatch(machine_number) and machine_number not in machine_map:
                        machine = db.Machine(
                            plant_id=plant.id,
                            area_id=area.id,
                            machine_number=machine_number,
                            machine_name=f"Machine {machine_number}",
                            status_id=_lookup(session, db.AssetStatus, "Active"),
                            source_system="legacy_excel",
                            source_import_batch_id=batch.id,
                        )
                        session.add(machine)
                        session.flush()
                        machine_map[machine_number] = machine
                    tool_number = _text(row.get("Tool #"))
                    if not _missing(tool_number) and tool_number not in tool_map:
                        tool = db.Tool(
                            business_identifier=tool_number,
                            tool_number=tool_number,
                            mold_number=tool_number,
                            display_name=f"Tool {tool_number}",
                            status_id=_lookup(session, db.AssetStatus, "Active"),
                            source_system="legacy_excel",
                            source_import_batch_id=batch.id,
                        )
                        session.add(tool)
                        session.flush()
                        tool_map[tool_number] = tool
                compatible_status = _lookup(
                    session, db.CompatibilityStatus, "Observed", display="Observed in legacy source"
                )
                compatibility_source = _lookup(
                    session, db.CompatibilitySource, "legacy_workbook", display="Legacy workbook"
                )
                effective = started.replace(tzinfo=None)
                em_pairs: set[tuple[int, int]] = set()
                et_pairs: set[tuple[int, int]] = set()
                tm_pairs: set[tuple[int, int]] = set()
                import_rows: dict[tuple[str, int], db.ImportRow] = {}
                audit_completed_type_id = session.scalar(
                    select(db.HistoryEventType.id).where(db.HistoryEventType.code == "audit_completed")
                )
                if audit_completed_type_id is None:
                    raise RuntimeError("Required history event type 'audit_completed' is unavailable.")
                history_event_count = 0
                for number, row in rows["EOAT Inventory"]:
                    eoat_id = _text(row.get("EOAT Assembly ID"))
                    machine_number = _text(row.get("Press/Machine #"))
                    tool_number = _text(row.get("Tool #"))
                    # A compatibility-only source row for a repeated legacy ID
                    # is not evidence for one arbitrary physical unit.  Preserve
                    # it as provenance until a design/family relationship can be
                    # made explicitly, rather than collapsing identities again.
                    compatibility_only = _text(row.get("Entry Type")).casefold() != "audited"
                    eoat = None if compatibility_only and eoat_id in split_source_identifiers else eoat_map.get(eoat_id)
                    machine = machine_map.get(machine_number)
                    tool = tool_map.get(tool_number)
                    audit_id = _text(row.get("Audit ID"))
                    audit = None
                    if not compatibility_only:
                        audit = db.AuditRecord(
                            audit_identifier=audit_id,
                            eoat_id=eoat.id if eoat else None,
                            machine_id=machine.id if machine else None,
                            tool_id=tool.id if tool else None,
                            audit_date=_datetime(row.get("Audit Date")),
                            status_id=_lookup(session, db.AssetStatus, row.get("Status")),
                            source_sheet="EOAT Inventory",
                            source_row_number=number,
                            details_json=_json_safe(row),
                            notes=_text(row.get("Notes")) or None,
                            source_system="legacy_excel",
                            source_import_batch_id=batch.id,
                        )
                        session.add(audit)
                        session.flush()
                    if eoat is not None and audit is not None:
                        session.add(
                            db.EntityHistoryEvent(
                                event_uuid=str(
                                    uuid5(NAMESPACE_URL, f"{before_checksum}|audit-history|{audit.id}|{audit_id}")
                                ),
                                entity_type="eoat",
                                entity_id=eoat.id,
                                event_type_id=audit_completed_type_id,
                                occurred_at=audit.audit_date or started,
                                event_category="AUDITS",
                                summary=f"Audit {audit_id}",
                                description="Documented legacy audit imported into EOAT Atlas.",
                                notes=audit.notes,
                                source_table="audit_records",
                                source_record_id=audit.id,
                                metadata_json={
                                    "audit_id": audit_id,
                                    "import_batch_uuid": result.import_batch_uuid,
                                    "source_sheet": "EOAT Inventory",
                                    "source_row": number,
                                    "import_provenance": "structured_audit_record",
                                },
                            )
                        )
                        history_event_count += 1
                    import_row = db.ImportRow(
                        import_batch_id=batch.id,
                        source_sheet="EOAT Inventory",
                        source_row_number=number,
                        source_identifier=audit_id,
                        target_entity_type="audit_record" if audit is not None else "compatibility_evidence",
                        target_entity_id=audit.id if audit is not None else None,
                        status="IMPORTED" if audit is not None else "EVIDENCE_ONLY",
                        raw_values_json=_json_safe(row),
                        normalized_values_json={
                            "eoat": eoat_id,
                            "machine": machine_number if machine else None,
                            "tool": tool_number if tool else None,
                        },
                    )
                    session.add(import_row)
                    session.flush()
                    import_rows[("EOAT Inventory", number)] = import_row
                    if eoat and machine and (eoat.id, machine.id) not in em_pairs:
                        session.add(
                            db.EOATMachineCompatibility(
                                eoat_id=eoat.id,
                                machine_id=machine.id,
                                compatibility_status_id=compatible_status,
                                verification_source_id=compatibility_source,
                                effective_from=effective,
                                reason="Observed association in legacy audit source",
                                source_system="legacy_excel",
                                source_import_batch_id=batch.id,
                            )
                        )
                        em_pairs.add((eoat.id, machine.id))
                    if eoat and tool and (eoat.id, tool.id) not in et_pairs:
                        session.add(
                            db.EOATToolCompatibility(
                                eoat_id=eoat.id,
                                tool_id=tool.id,
                                compatibility_status_id=compatible_status,
                                verification_source_id=compatibility_source,
                                effective_from=effective,
                                reason="Observed association in legacy audit source",
                                source_system="legacy_excel",
                                source_import_batch_id=batch.id,
                            )
                        )
                        et_pairs.add((eoat.id, tool.id))
                    if tool and machine and (tool.id, machine.id) not in tm_pairs:
                        session.add(
                            db.ToolMachineCompatibility(
                                tool_id=tool.id,
                                machine_id=machine.id,
                                compatibility_status_id=compatible_status,
                                verification_source_id=compatibility_source,
                                effective_from=effective,
                                reason="Observed association in legacy audit source",
                                source_system="legacy_excel",
                                source_import_batch_id=batch.id,
                            )
                        )
                        tm_pairs.add((tool.id, machine.id))
                photo_type = _lookup(session, db.DocumentType, "photo", display="Photo")
                project_root = source.parents[2]
                valid_photo_count = 0
                for number, row in rows["Photo Index"]:
                    photo_id = _text(row.get("Photo ID"))
                    relative = _text(row.get("Stored Relative Path"))
                    folder = _text(row.get("Folder Path"))
                    filename = _text(row.get("Stored Filename")) or _text(row.get("Photo Filename"))
                    candidate = project_root / relative if relative else project_root / folder / filename
                    raw = _json_safe(row)
                    if not relative and not (folder and filename):
                        import_row = db.ImportRow(
                            import_batch_id=batch.id,
                            source_sheet="Photo Index",
                            source_row_number=number,
                            source_identifier=photo_id,
                            target_entity_type="photo",
                            status="DEFERRED",
                            raw_values_json=raw,
                            error_summary="Placeholder row has no valid path",
                        )
                        session.add(import_row)
                        session.flush()
                        import_rows[("Photo Index", number)] = import_row
                        continue
                    document = db.Document(
                        document_uuid=str(uuid5(NAMESPACE_URL, f"{before_checksum}|photo|{number}|{photo_id}")),
                        document_type_id=photo_type,
                        document_number=photo_id or None,
                        title=_text(row.get("Description")) or filename,
                        description=_text(row.get("Notes")) or None,
                        file_name=filename,
                        file_extension=Path(filename).suffix or None,
                        storage_path=str(candidate),
                        mime_type=mimetypes.guess_type(filename)[0],
                        source_system="legacy_excel",
                        source_import_batch_id=batch.id,
                    )
                    session.add(document)
                    session.flush()
                    photo = db.Photo(
                        document_id=document.id,
                        photo_view_type=_text(row.get("Photo Type")) or _text(row.get("EOAT Area Shown")) or None,
                        captured_at=_datetime(row.get("Date Taken")),
                        caption=_text(row.get("Description")) or None,
                    )
                    session.add(photo)
                    session.flush()
                    valid_photo_count += 1
                    for entity_type, entity_identifier, entity in (
                        ("eoat", _text(row.get("EOAT Assembly ID")), eoat_map.get(_text(row.get("EOAT Assembly ID")))),
                        ("tool", _text(row.get("Tool #")), tool_map.get(_text(row.get("Tool #")))),
                        (
                            "machine",
                            _text(row.get("Press/Machine #")),
                            machine_map.get(_text(row.get("Press/Machine #"))),
                        ),
                    ):
                        if entity:
                            session.add(
                                db.DocumentLink(
                                    document_id=document.id,
                                    entity_type=entity_type,
                                    entity_id=entity.id,
                                    relationship_type="legacy_photo",
                                    is_primary=False,
                                )
                            )
                    import_row = db.ImportRow(
                        import_batch_id=batch.id,
                        source_sheet="Photo Index",
                        source_row_number=number,
                        source_identifier=photo_id,
                        target_entity_type="photo",
                        target_entity_id=photo.id,
                        status="IMPORTED",
                        raw_values_json=raw,
                        normalized_values_json={"storage_path": str(candidate)},
                    )
                    session.add(import_row)
                    session.flush()
                    import_rows[("Photo Index", number)] = import_row
                for item in review_items:
                    related_row = import_rows.get((item.source_sheet, item.source_row or -1))
                    session.add(
                        db.ImportIssue(
                            import_batch_id=batch.id,
                            import_row_id=related_row.id if related_row else None,
                            severity=item.severity,
                            issue_code=item.issue_code,
                            field_name=item.field_name or None,
                            source_value=item.source_value or None,
                            description=item.current_proposed_action,
                            suggested_resolution="; ".join(item.allowed_resolution_actions),
                            resolution_notes=json.dumps(
                                {
                                    "issue_id": item.issue_id,
                                    "status": item.resolution_status,
                                    "related_source_rows": item.related_source_rows,
                                }
                            ),
                        )
                    )
                counts = {
                    "plants": 1,
                    "areas": len(area_map),
                    "eoats": len(eoat_map),
                    "machines": len(machine_map),
                    "tools": len(tool_map),
                    "parts": 0,
                    "eoat_machine_compatibility": len(em_pairs),
                    "eoat_tool_compatibility": len(et_pairs),
                    "tool_machine_compatibility": len(tm_pairs),
                    "audit_records": sum(
                        _text(row.get("Entry Type")).casefold() == "audited"
                        for _, row in rows["EOAT Inventory"]
                    ),
                    "entity_history_events": history_event_count,
                    "documents": valid_photo_count,
                    "photos": valid_photo_count,
                    "import_rows": sum(result.rows_discovered_by_sheet.values()),
                    "import_issues": len(review_items),
                    "eoat_installations": 0,
                }
                batch.status = "COMPLETED"
                batch.completed_at = datetime.now(timezone.utc)
                batch.records_imported = sum(counts.values())
                batch.records_rejected = 0
                batch.warnings_count = len(review_items)
                result.records_imported_by_table = counts
                result.transaction_result = "COMMITTED"
            result.foreign_key_validation = (
                "PASS"
                if session.scalar(
                    text(
                        "SELECT COUNT(*) FROM information_schema.referential_constraints WHERE constraint_schema = DATABASE()"
                    )
                )
                else "PASS_NO_FOREIGN_KEYS_REPORTED"
            )
    result.warnings = len(review_items)
    result.unresolved_issues = sum(item.resolution_status in {"UNRESOLVED", "DEFERRED"} for item in review_items)
    result.resolved_issues = sum(
        item.resolution_status in {"RESOLVED", "NOT_APPLICABLE", "REJECTED_WITH_REASON"} for item in review_items
    )
    result.missing_relationships = sum(
        item.issue_code
        in {
            "MISSING_MACHINE",
            "MISSING_TOOL",
            "AMBIGUOUS_MACHINE_VALUE",
            "CONFLICTING_CURRENT_ASSIGNMENT",
            "CURRENT_LOCATION_UNKNOWN",
        }
        for item in review_items
    )
    result.deferred_parts = sum(item.issue_code == "POSSIBLE_PART_NOT_CONFIRMED" for item in review_items)
    result.deferred_installations = sum(item.issue_code == "INSTALLATION_DATE_UNKNOWN" for item in review_items)
    paths = []
    for _, row in rows["Photo Index"]:
        relative = _text(row.get("Stored Relative Path"))
        folder = _text(row.get("Folder Path"))
        filename = _text(row.get("Stored Filename")) or _text(row.get("Photo Filename"))
        if relative or (folder and filename):
            paths.append(source.parents[2] / relative if relative else source.parents[2] / folder / filename)
    result.photo_path_validation = {
        "checked": len(paths),
        "existing": sum(path.exists() for path in paths),
        "missing": sum(not path.exists() for path in paths),
    }
    result.document_path_validation = dict(result.photo_path_validation)
    result.completed_at = datetime.now(timezone.utc).isoformat()
    after_stat = source.stat()
    result.source_workbook_unchanged = (
        before_stat.st_size == after_stat.st_size
        and before_stat.st_mtime_ns == after_stat.st_mtime_ns
        and before_checksum == _checksum(source)
    )
    write_import_report(result, report_output)
    return result


def write_import_report(result: ImportResult, output_directory: str | Path) -> tuple[Path, Path]:
    output = Path(output_directory)
    output.mkdir(parents=True, exist_ok=True)
    json_path = output / "development_import_report.json"
    md_path = output / "development_import_report.md"
    json_path.write_text(json.dumps(result.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")
    counts = (
        "\n".join(f"| {table} | {count} |" for table, count in sorted(result.records_imported_by_table.items()))
        or "| - | 0 |"
    )
    md_path.write_text(
        "# EOAT Atlas Development MySQL Import\n\n"
        f"- Source: `{result.source_workbook}`\n- SHA-256: `{result.source_checksum}`\n- Batch UUID: `{result.import_batch_uuid}`\n"
        f"- Schema revision: `{result.schema_revision}`\n- Started: `{result.started_at}`\n- Completed: `{result.completed_at}`\n"
        f"- Transaction: **{result.transaction_result}**\n- Source unchanged: **{result.source_workbook_unchanged}**\n"
        f"- Warnings/issues: **{result.warnings}**; unresolved/deferred: **{result.unresolved_issues}**; rejected: **{result.records_rejected}**\n"
        f"- Deferred parts/installations: **{result.deferred_parts} / {result.deferred_installations}**\n\n"
        "## Imported records\n\n| Table | Count |\n|---|---:|\n" + counts + "\n",
        encoding="utf-8",
    )
    return md_path, json_path


def database_counts() -> dict[str, int]:
    factory = create_session_factory()
    names = [
        "plants",
        "areas",
        "eoats",
        "machines",
        "tools",
        "parts",
        "eoat_machine_compatibility",
        "eoat_tool_compatibility",
        "tool_machine_compatibility",
        "audit_records",
        "documents",
        "photos",
        "import_rows",
        "import_issues",
        "eoat_installations",
    ]
    with factory() as session:
        return {name: int(session.scalar(text(f"SELECT COUNT(*) FROM {name}")) or 0) for name in names}
